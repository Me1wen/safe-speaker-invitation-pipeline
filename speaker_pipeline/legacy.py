"""Safe migration from the three original packages' 11/12-column tables."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .common import candidate_id, clean, read_csv, today_iso

LEGACY_COLUMNS = [
    "Highest Level",
    "Award Year",
    "Full Name",
    "University/Institute",
    "Research Field",
    "Placeholder",
    "Title and Last Name",
    "Email",
    "Phone",
    "Link",
    "Notes",
]

CONTACT_LABELS = {
    "corporate communications": "Corporate Communications",
    "media relations": "Media Relations",
    "investor relations": "Investor Relations",
    "public affairs": "Public Affairs",
    "executive office": "Executive Office",
    "university relations": "University Relations",
    "speaker inquiry": "Speaker Inquiry",
    "general contact": "General Contact",
}


def read_legacy_table(path: Path, sheet_name: str | None = None) -> list[dict[str, str]]:
    if path.suffix.casefold() == ".csv":
        rows = read_csv(path)
        headers = list(rows[0]) if rows else []
        if not rows:
            with path.open("r", encoding="utf-8-sig") as handle:
                headers = [clean(value) for value in handle.readline().rstrip("\r\n").split(",")]
    elif path.suffix.casefold() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = (
                workbook[sheet_name]
                if sheet_name
                else (
                    workbook["Draft Results"]
                    if "Draft Results" in workbook.sheetnames
                    else workbook.active
                )
            )
            values = worksheet.iter_rows(values_only=True)
            headers = [clean(value) for value in next(values, ())]
            rows = [
                {header: clean(value) for header, value in zip(headers, values_row, strict=False)}
                for values_row in values
                if any(clean(value) for value in values_row)
            ]
        finally:
            workbook.close()
    else:
        raise ValueError("Legacy input must be CSV, XLSX, or XLSM.")
    missing = [column for column in LEGACY_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Legacy table is missing: {', '.join(missing)}")
    return rows


def _email_parts(value: str) -> tuple[str, str]:
    value = clean(value)
    matches = re.findall(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", value, flags=re.I)
    if len(matches) != 1:
        return "", ""
    email = matches[0].casefold()
    prefix = clean(value[: value.casefold().find(email.casefold())].strip(" -:;"))
    contact_type = CONTACT_LABELS.get(
        prefix.casefold(), "Direct" if not prefix else "General Contact"
    )
    return email, contact_type


def migrate_legacy_rows(rows: list[dict[str, str]], speaker_type: str) -> list[dict[str, str]]:
    normalized_type = speaker_type.casefold()
    if normalized_type not in {"academic", "industry"}:
        raise ValueError("speaker_type must be academic or industry")
    output: list[dict[str, str]] = []
    display_type = normalized_type.title()
    for raw in rows:
        name = clean(raw.get("Full Name"))
        organization = clean(raw.get("University/Institute"))
        email, contact_type = _email_parts(raw.get("Email", ""))
        profile_url = clean(raw.get("Link"))
        salutation = clean(raw.get("Title and Last Name")).rstrip(",")
        if normalized_type == "industry" and re.fullmatch(
            r".+,", clean(raw.get("Title and Last Name"))
        ):
            salutation = name
        notes = clean(raw.get("Notes"))
        legacy_status = clean(raw.get("Review Status"))
        if legacy_status == "Approved":
            notes = (
                f"{notes} | Legacy Approved; reapproval required after migration"
                if notes
                else "Legacy Approved; reapproval required after migration"
            )
        output.append(
            {
                "Candidate ID": candidate_id(display_type, name, organization)
                if name and organization
                else "",
                "Speaker Type": display_type,
                "Review Status": "Needs Review",
                "Full Name": name,
                "Organization": organization,
                "Title": clean(raw.get("Highest Level")),
                "Expertise": clean(raw.get("Research Field")),
                "Topic Fit": "",
                "Preferred Salutation": salutation,
                "Email": email,
                "Contact Type": contact_type,
                "Email Source URL": profile_url if contact_type == "Direct" else "",
                "Phone": clean(raw.get("Phone")),
                "Profile URL": profile_url,
                "Discovery Source URL": "",
                "Notes": notes,
                "Last Checked": today_iso(),
            }
        )
    return output
