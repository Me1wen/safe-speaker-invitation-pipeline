"""Shared table, identity, URL, and merge helpers."""

from __future__ import annotations

import csv
import hashlib
import os
import re
import tempfile
from collections.abc import Iterable
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .schema import CANDIDATE_COLUMNS, CANDIDATE_STATUSES

EMAIL_RE = re.compile(r"^[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}$", re.IGNORECASE)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


MULTILINE_COLUMNS = {"Body Text", "Evidence Summary", "Snippet"}


def normalize_multiline(value: object) -> str:
    """Normalize line endings without destroying intentional paragraph breaks."""
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def clean_cell(column: str, value: object) -> str:
    return normalize_multiline(value) if column in MULTILINE_COLUMNS else clean(value)


def today_iso() -> str:
    return date.today().isoformat()


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalize_name(value: str) -> str:
    value = re.sub(r"\b(?:dr|prof|professor|mr|mrs|ms)\.?\b", "", clean(value), flags=re.I)
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def normalize_url(value: str) -> str:
    parsed = urlparse(clean(value))
    if not parsed.scheme or not parsed.netloc:
        return clean(value)
    path = re.sub(r"/+", "/", parsed.path).rstrip("/")
    host = parsed.netloc.casefold().removeprefix("www.")
    return f"{parsed.scheme.casefold()}://{host}{path}"


def candidate_id(speaker_type: str, name: str, organization: str) -> str:
    normalized = (
        f"{speaker_type.casefold()}|{normalize_name(name)}|{clean(organization).casefold()}"
    )
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:12].upper()
    prefix = "ACAD" if speaker_type.casefold() == "academic" else "IND"
    return f"{prefix}-{digest}"


def draft_id(candidate_id_value: str, subject: str, body: str, campaign_id: str = "") -> str:
    digest = hashlib.sha256(
        f"{campaign_id}|{candidate_id_value}|{subject}|{normalize_multiline(body)}".encode()
    ).hexdigest()
    return f"DRAFT-{digest[:16].upper()}"


def split_domains(value: str, fallback: str = "") -> tuple[str, ...]:
    values = [*re.split(r"[;,\s]+", clean(value)), fallback]
    domains: set[str] = set()
    for item in values:
        item = clean(item).casefold()
        if not item:
            continue
        parsed = urlparse(item if "://" in item else f"https://{item}")
        host = parsed.hostname or ""
        if host:
            domains.add(host.removeprefix("www."))
    return tuple(sorted(domains))


def host_allowed(url: str, domains: Iterable[str]) -> bool:
    parsed = urlparse(clean(url))
    host = (parsed.hostname or "").casefold().removeprefix("www.")
    return parsed.scheme in {"http", "https"} and any(
        host == domain or host.endswith(f".{domain}") for domain in domains
    )


def email_allowed(email: str, domains: Iterable[str]) -> bool:
    if not EMAIL_RE.fullmatch(clean(email)):
        return False
    host = clean(email).rsplit("@", 1)[-1].casefold().removeprefix("www.")
    return any(host == domain or host.endswith(f".{domain}") for domain in domains)


def _compatible_headers(headers: list[str], expected: list[str]) -> bool:
    """Accept only the exact current schema or an explicitly known legacy schema."""
    if headers == expected:
        return True
    from . import schema

    legacy_candidates = getattr(schema, "V1_0_CANDIDATE_COLUMNS", [])
    legacy_drafts = getattr(schema, "V1_0_DRAFT_COLUMNS", [])
    legacy_v11_drafts = getattr(schema, "DRAFT_COLUMNS_V1_1_LEGACY", [])
    return (expected == schema.CANDIDATE_COLUMNS and headers == legacy_candidates) or (
        expected == schema.DRAFT_COLUMNS
        and tuple(headers) in {tuple(legacy_drafts), tuple(legacy_v11_drafts)}
    )


def _legacy_schema_version(headers: list[str], expected: list[str]) -> str:
    from . import schema

    if expected == schema.DRAFT_COLUMNS and headers == getattr(
        schema, "DRAFT_COLUMNS_V1_1_LEGACY", []
    ):
        return "1.1"
    return "1.0"


def read_csv(path: Path, expected: list[str] | None = None) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        if expected is not None and not _compatible_headers(headers, expected):
            raise ValueError(f"{path} columns do not match the required schema. Found: {headers}")
        columns = expected or headers
        rows = [
            {column: clean_cell(column, row.get(column, "")) for column in columns}
            for row in reader
        ]
        if expected is not None and headers != expected and "Schema Version" in columns:
            for row in rows:
                row["Schema Version"] = _legacy_schema_version(headers, expected)
        return rows


def write_csv(path: Path, columns: list[str], rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({column: row.get(column, "") for column in columns})
            handle.flush()
            os.fsync(handle.fileno())
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def read_table(
    path: Path, columns: list[str], sheet_name: str | None = None
) -> list[dict[str, str]]:
    if path.suffix.casefold() == ".csv":
        return read_csv(path, columns)
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("Input must be CSV, XLSX, or XLSM.")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Sheet not found: {sheet_name}")
            worksheet = workbook[sheet_name]
        elif len(workbook.sheetnames) == 1:
            worksheet = workbook.active
        else:
            raise ValueError("Workbook has multiple sheets; provide a sheet name.")
        cell_rows = worksheet.iter_rows()
        headers = [clean(cell.value) for cell in next(cell_rows, ())]
        if not _compatible_headers(headers, columns):
            raise ValueError(f"Workbook columns do not match the required schema. Found: {headers}")
        rows: list[dict[str, str]] = []
        for worksheet_row in cell_rows:
            for cell in worksheet_row:
                if cell.data_type == "f":
                    raise ValueError(
                        f"Workbook formulas are not allowed in review tables ({cell.coordinate})."
                    )
            values_by_header = {
                header: cell.value for header, cell in zip(headers, worksheet_row, strict=False)
            }
            row = {
                column: clean_cell(column, values_by_header.get(column, "")) for column in columns
            }
            if any(row.values()):
                if headers != columns and "Schema Version" in columns:
                    row["Schema Version"] = _legacy_schema_version(headers, columns)
                rows.append(row)
        return rows
    finally:
        workbook.close()


def write_workbook(
    path: Path, sheet_name: str, columns: list[str], rows: Iterable[dict[str, object]]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.append(columns)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_number, row in enumerate(rows, start=2):
        for column_number, column in enumerate(columns, start=1):
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.value = str(row.get(column, "") or "")
            # openpyxl otherwise serializes leading '=' as an executable formula.
            cell.data_type = "s"
            if column in MULTILINE_COLUMNS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column_cells in worksheet.columns:
        width = max(len(clean(cell.value)) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(width + 2, 12), 60
        )
    worksheet.auto_filter.ref = worksheet.dimensions
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=path.suffix, dir=path.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def candidate_key(row: dict[str, str]) -> tuple[str, str, str]:
    return (
        clean(row.get("Speaker Type")).casefold(),
        normalize_name(row.get("Full Name", "")),
        clean(row.get("Organization")).casefold(),
    )


def merge_candidates(
    existing: Iterable[dict[str, str]],
    discovered: Iterable[dict[str, str]],
) -> list[dict[str, str]]:
    """Merge without overwriting any manually approved record."""
    merged: list[dict[str, str]] = []
    by_key: dict[tuple[str, str, str], int] = {}
    by_url: dict[str, int] = {}

    for source_row in existing:
        row = {column: clean(source_row.get(column, "")) for column in CANDIDATE_COLUMNS}
        if row["Review Status"] not in CANDIDATE_STATUSES:
            raise ValueError(f"Invalid Review Status: {row['Review Status']}")
        index = len(merged)
        merged.append(row)
        key = candidate_key(row)
        if all(key):
            by_key[key] = index
        if row["Profile URL"]:
            by_url[normalize_url(row["Profile URL"])] = index

    manual_fields = {"Review Status", "Topic Fit", "Preferred Salutation", "Notes"}
    for source_row in discovered:
        row = {column: clean(source_row.get(column, "")) for column in CANDIDATE_COLUMNS}
        key = candidate_key(row)
        url_key = normalize_url(row["Profile URL"])
        index = by_key.get(key) if all(key) else None
        if index is None and url_key:
            index = by_url.get(url_key)
        if index is None:
            index = len(merged)
            merged.append(row)
        else:
            old = merged[index]
            if old.get("Review Status") != "Approved":
                combined = dict(old)
                for column in CANDIDATE_COLUMNS:
                    if column not in manual_fields and row.get(column):
                        combined[column] = row[column]
                for column in manual_fields:
                    combined[column] = old.get(column) or row.get(column, "")
                merged[index] = combined
        final = merged[index]
        final["Candidate ID"] = final.get("Candidate ID") or candidate_id(
            final.get("Speaker Type", ""), final.get("Full Name", ""), final.get("Organization", "")
        )
        final["Review Status"] = final.get("Review Status") or "Draft"
        final["Last Checked"] = final.get("Last Checked") or today_iso()
        final["Schema Version"] = final.get("Schema Version") or "1.1"
        final["Revision"] = final.get("Revision") or "1"
        final_key = candidate_key(final)
        if all(final_key):
            by_key[final_key] = index
        if final.get("Profile URL"):
            by_url[normalize_url(final["Profile URL"])] = index
    return merged
