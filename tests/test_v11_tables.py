import pytest
from openpyxl import Workbook, load_workbook

from speaker_pipeline.common import read_csv, read_table, write_csv, write_workbook
from speaker_pipeline.schema import DRAFT_COLUMNS, V1_0_DRAFT_COLUMNS


def _draft_row():
    row = {column: "" for column in DRAFT_COLUMNS}
    row.update(
        {
            "Draft ID": "DRAFT-1",
            "Candidate ID": "IND-1",
            "Draft Status": "Draft",
            "Subject": "A safe invitation",
            "Body Text": "First paragraph.\n\nSecond paragraph.",
        }
    )
    return row


def test_body_newlines_survive_csv_and_xlsx(tmp_path):
    row = _draft_row()
    csv_path = tmp_path / "drafts.csv"
    xlsx_path = tmp_path / "drafts.xlsx"
    write_csv(csv_path, DRAFT_COLUMNS, [row])
    write_workbook(xlsx_path, "Email Drafts", DRAFT_COLUMNS, [row])
    assert read_csv(csv_path, DRAFT_COLUMNS)[0]["Body Text"] == row["Body Text"]
    assert read_table(xlsx_path, DRAFT_COLUMNS)[0]["Body Text"] == row["Body Text"]


def test_workbook_formula_like_values_are_literal_strings(tmp_path):
    row = _draft_row()
    row["Subject"] = '=HYPERLINK("https://example.test")'
    path = tmp_path / "review.xlsx"
    write_workbook(path, "Email Drafts", DRAFT_COLUMNS, [row])
    workbook = load_workbook(path, data_only=False)
    try:
        cell = workbook["Email Drafts"].cell(2, DRAFT_COLUMNS.index("Subject") + 1)
        assert cell.data_type == "s"
        assert cell.value.startswith("=")
    finally:
        workbook.close()


def test_workbook_import_rejects_formula_cells(tmp_path):
    path = tmp_path / "unsafe.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(DRAFT_COLUMNS)
    sheet.append(["=1+1", *([""] * (len(DRAFT_COLUMNS) - 1))])
    workbook.save(path)
    with pytest.raises(ValueError, match="formulas are not allowed"):
        read_table(path, DRAFT_COLUMNS)


def test_exact_v1_table_upgrades_but_unknown_header_is_rejected(tmp_path):
    legacy = tmp_path / "legacy.csv"
    write_csv(legacy, V1_0_DRAFT_COLUMNS, [{column: "" for column in V1_0_DRAFT_COLUMNS}])
    rows = read_csv(legacy, DRAFT_COLUMNS)
    assert set(rows[0]) == set(DRAFT_COLUMNS)
    assert rows[0].get("Approval Hash", "") == ""

    bad = tmp_path / "bad.csv"
    write_csv(bad, [*V1_0_DRAFT_COLUMNS, "Unexpected"], [])
    with pytest.raises(ValueError, match="required schema"):
        read_csv(bad, DRAFT_COLUMNS)
